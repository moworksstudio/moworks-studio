"""Moworks Studio intake: makes the user's files readable for Claude before the Assets analysis.

Written for the app from the ideas of the Moworks Motion kit 2.0.4 (engine/intake/psd.py, commit 30b7dcb, MIT, Moworks): the layer tree,
frames from artboards or top-level groups, text layers with their fonts, layer export for keep-design films. No kit code runs here.

  python intake.py psd <file.psd|psb> <out> [--width 1600]          tree, composite, one picture per artboard / group frame, texts, fonts
  python intake.py psd-export <file> <out> --frame <name> [--width 1920]   one frame's visible layers as PNGs + layout.json (keep design)
  python intake.py pptx <file.pptx> <out>                           slides: texts, notes, layout, pictures (slide pictures: PowerPoint, by the app)
  python intake.py pdf <file.pdf|ai> <out> [--width 1200] [--logo]  the whole document (pypdfium2), for a film's files and a Brand Kit alike: every page's text in
                                                                    TEXT-NN.md, every page as a picture (pages/) and on numbered contact sheets (sheet-NN.jpg);
                                                                    --logo: page 1 on no ground (logo.png)
  python intake.py sheet <list.txt> <out>                           pictures named one a line, on contact sheets (sheet-NN.jpg)
  python intake.py music <file> <out>                               length, tempo, beats, the loudest passages (librosa)

Libraries (pinned in config/pins.json python.intake, installed as the app's components): psd-tools (MIT), python-pptx (MIT),
pypdfium2 (Apache-2.0 / BSD-3), librosa (ISC), soundfile (BSD-3), Pillow (MIT-CMU), numpy (BSD-3).
Prints one JSON object on stdout; writes its files and a readable summary (*.md) into <out>."""
import json, os, re, sys
from pathlib import Path

for _s in (sys.stdout, sys.stderr):
    try: _s.reconfigure(encoding="utf-8", errors="replace")
    except Exception: pass
_libs = os.environ.get("MW_PYLIBS")          # the app's own component folder (runtime/python-libs), first on the path
if _libs and os.path.isdir(_libs): sys.path.insert(0, _libs)


def slug(s): return re.sub(r"[^A-Za-z0-9]+", "-", str(s)).strip("-")[:40] or "x"
def arg(name, default=None):
    a = sys.argv
    return a[a.index(name) + 1] if name in a and a.index(name) + 1 < len(a) else default
def done(obj): print(json.dumps(obj, ensure_ascii=False)); sys.exit(0)
def fail(msg): print(json.dumps({"ok": False, "error": msg}, ensure_ascii=False)); sys.exit(0)


# ---- PSD / PSB -----------------------------------------------------------------------------------------------------------------------------
def _bbox(l):
    try: x1, y1, x2, y2 = l.bbox; return [int(x1), int(y1), int(x2 - x1), int(y2 - y1)]
    except Exception: return [0, 0, 0, 0]

def _text(l):
    """A type layer: its words, fonts, size (pt scaled by the layer's transform) and colour, as far as the file says."""
    out = {"text": (getattr(l, "text", "") or "").replace("\r", "\n").strip()}
    try:
        fonts = [f.get("Name", "") for f in l.resource_dict["FontSet"]]
        runs = l.engine_dict["StyleRun"]["RunArray"]
        sheet = runs[0]["StyleSheet"]["StyleSheetData"]
        scale = abs(float(l.transform[3])) if getattr(l, "transform", None) else 1.0
        out["font"] = str(fonts[int(sheet.get("Font", 0))]).strip("'") if fonts else None
        if "FontSize" in sheet: out["size"] = round(float(sheet["FontSize"]) * scale, 1)
        fill = sheet.get("FillColor", {}).get("Values")
        if fill and len(fill) >= 4: out["color"] = "#%02X%02X%02X" % tuple(int(round(float(v) * 255)) for v in fill[1:4])
        out["fonts"] = sorted({str(fonts[int(r["StyleSheet"]["StyleSheetData"].get("Font", 0))]).strip("'") for r in runs if fonts})
    except Exception: pass
    return out

def _save(img, path, width):
    if img is None: return None
    if width and img.width > width: img = img.resize((width, max(1, round(img.height * width / img.width))))
    img.save(path)
    return path.name

def psd_scan(file, out, width):
    from psd_tools import PSDImage
    psd = PSDImage.open(file)
    tree, texts, fonts = [], [], set()
    def walk(group, depth, path):
        for l in group:
            p = f"{path}/{l.name}" if path else l.name
            node = {"name": l.name, "path": p, "kind": l.kind, "visible": l.is_visible(), "bbox": _bbox(l), "depth": depth}
            if l.kind == "type":
                t = _text(l); texts.append({"layer": p, **t}); fonts.update(t.get("fonts") or ([t["font"]] if t.get("font") else []))
            if l.is_group(): node["layers"] = sum(1 for _ in l.descendants())
            tree.append(node)
            if l.is_group() and depth < 3: walk(l, depth + 1, p)
    walk(psd, 0, "")
    top = [l for l in psd]
    boards = [l for l in top if l.kind == "artboard"]
    groups = [l for l in top if l.is_group() and l.kind != "artboard"]
    frames = boards or (groups if 2 <= len(groups) <= 30 else [])
    composite = None
    try: composite = _save(psd.topil() or psd.composite(), out / "composite.png", width)
    except Exception:
        try: composite = _save(psd.composite(), out / "composite.png", width)
        except Exception: composite = None
    pictures = []
    for i, f in enumerate(frames[:24]):
        name = f"frame-{i + 1:02d}-{slug(f.name)}.png"
        box = _bbox(f)
        try:
            was = f.visible; f.visible = True           # a hidden frame still gets its picture (the kit's trap: hidden ancestors draw nothing)
            img = f.composite()
            if img is not None and box[2] <= 0: box = [0, 0, img.width, img.height]          # a hidden group reports an empty box: the picture's size
            pic = _save(img, out / name, width); f.visible = was
        except Exception: pic = None
        pictures.append({"frame": f.name, "kind": f.kind, "bbox": box, "picture": pic, "hidden": not f.is_visible()})
    info = {"ok": True, "kind": "psd", "width": psd.width, "height": psd.height, "color_mode": getattr(psd.color_mode, "name", str(psd.color_mode)),
            "artboards": len(boards), "frames": pictures, "composite": composite, "texts": texts[:200], "fonts": sorted(f for f in fonts if f),
            "tree": tree[:400]}
    md = [f"# {Path(file).name}", "", f"{psd.width} x {psd.height} px, {info['color_mode']}, {len(boards)} artboard(s), {len(tree)} layers listed (to depth 3).", ""]
    if composite: md += [f"Whole file: `{composite}`", ""]
    if pictures:
        md += ["## Frames (" + ("artboards" if boards else "top-level groups") + ")", ""]
        md += [f"- {p['frame']} ({p['bbox'][2]}x{p['bbox'][3]} at {p['bbox'][0]},{p['bbox'][1]}{', hidden in the file' if p.get('hidden') else ''}): `{p['picture'] or 'no picture'}`" for p in pictures]
        md.append("")
    md += ["## Fonts", "", ", ".join(info["fonts"]) or "none found", "", "## Texts", ""]
    md += [f"- {t['layer']}: \"{t['text'][:160]}\"" + (" (" + ", ".join(str(v) for v in [t.get("font"), f"{t['size']} pt" if t.get("size") else None, t.get("color")] if v) + ")" if t.get("font") else "") for t in texts[:80]]
    md += ["", "## Layer tree", ""]
    md += [f"{'  ' * n['depth']}- {n['name']} [{n['kind']}{'' if n['visible'] else ', hidden'}] {n['bbox'][2]}x{n['bbox'][3]}" + (f" ({n['layers']} layers)" if n.get("layers") else "") for n in tree[:300]]
    (out / "PSD.md").write_text("\n".join(md) + "\n", encoding="utf-8")
    (out / "psd.json").write_text(json.dumps(info, ensure_ascii=False, indent=1), encoding="utf-8")
    done({**{k: info[k] for k in ("ok", "kind", "width", "height", "artboards", "composite", "fonts")}, "frames": len(pictures), "texts": len(texts), "md": "PSD.md"})

def psd_export(file, out, frame_name, width):
    """One frame's visible layers, back to front, each as a PNG at film scale, with layout.json: what a keep-design rebuild needs."""
    from psd_tools import PSDImage
    psd = PSDImage.open(file)
    frame = next((l for l in psd.descendants() if l.name == frame_name or (l.is_group() and frame_name.endswith("/" + l.name))), None)
    if frame is None and frame_name in ("", "*", "document"): frame = psd
    if frame is None: fail(f"no layer or group named {frame_name!r}")
    fx, fy, fw, fh = _bbox(frame) if frame is not psd else [0, 0, psd.width, psd.height]
    scale = (width / fw) if width and fw else 1.0
    units, n = [], 0
    def emit(l):
        nonlocal n
        x, y, w, h = _bbox(l)
        if w <= 0 or h <= 0: return
        n += 1
        name = f"{n:02d}_{slug(l.name)}.png"
        try: img = l.composite()
        except Exception: img = None
        if img is None: return
        if scale != 1.0: img = img.resize((max(1, round(img.width * scale)), max(1, round(img.height * scale))))
        img.save(out / name)
        unit = {"file": name, "name": l.name, "kind": l.kind, "x": round((x - fx) * scale, 1), "y": round((y - fy) * scale, 1), "w": img.width, "h": img.height,
                "opacity": round(l.opacity / 255, 3), "blend": str(l.blend_mode).split(".")[-1].lower()}
        if l.kind == "type":
            unit["text"] = _text(l)
            if unit["text"].get("size"): unit["text"]["size_px"] = round(unit["text"]["size"] * scale, 1)          # the size at the film's scale
        units.append(unit)
    def walk(g):
        for l in g:
            if not l.is_visible(): continue
            if l.kind == "adjustment": continue
            # a group with clipped layers (a photo in a shape) stays whole; other groups split down to their layers (the kit's rule)
            if l.is_group() and not any(getattr(c, "clipping", False) for c in l): walk(l)
            else: emit(l)
    was = getattr(frame, "visible", True)
    if frame is not psd: frame.visible = True
    walk(frame)
    if frame is not psd: frame.visible = was
    layout = {"frame": frame_name, "width": round(fw * scale), "height": round(fh * scale), "scale": round(scale, 4), "units": units}
    (out / "layout.json").write_text(json.dumps(layout, ensure_ascii=False, indent=1), encoding="utf-8")
    done({"ok": True, "kind": "psd-export", "frame": frame_name, "units": len(units), "layout": "layout.json", "width": layout["width"], "height": layout["height"]})


# ---- PowerPoint ----------------------------------------------------------------------------------------------------------------------------
def pptx_scan(file, out):
    from pptx import Presentation
    prs = Presentation(file)
    emu = lambda v: round(int(v or 0) / 9525)
    slides = []
    for i, s in enumerate(prs.slides, 1):
        texts, pics = [], []
        for sh in s.shapes:
            box = [emu(sh.left), emu(sh.top), emu(sh.width), emu(sh.height)]
            if getattr(sh, "has_text_frame", False) and sh.text_frame.text.strip():
                size = None
                try:
                    runs = [r for p in sh.text_frame.paragraphs for r in p.runs if r.font.size]
                    size = round(runs[0].font.size.pt, 1) if runs else None
                except Exception: pass
                texts.append({"text": sh.text_frame.text.strip()[:600], "box": box, "size": size, "name": sh.name})
            if sh.shape_type == 13:          # a picture
                try:
                    name = f"slide-{i:02d}-pic-{len(pics) + 1}.{sh.image.ext}"
                    (out / name).write_bytes(sh.image.blob)
                    pics.append({"file": name, "box": box})
                except Exception: pass
        notes = ""
        try:
            if s.has_notes_slide: notes = s.notes_slide.notes_text_frame.text.strip()
        except Exception: pass
        slides.append({"slide": i, "layout": s.slide_layout.name, "texts": texts, "pictures": pics, "notes": notes[:1500]})
    w, h = emu(prs.slide_width), emu(prs.slide_height)
    md = [f"# {Path(file).name}", "", f"{len(slides)} slide(s), {w} x {h} px (at 96 dpi).", ""]
    for s in slides:
        md += [f"## Slide {s['slide']} ({s['layout']})", ""]
        md += [f"- \"{t['text'][:300]}\"" + (f" ({t['size']} pt)" if t["size"] else "") for t in s["texts"]]
        md += [f"- picture: `{p['file']}`" for p in s["pictures"]]
        if s["notes"]: md += ["", f"Notes: {s['notes'][:600]}"]
        md.append("")
    (out / "PPTX.md").write_text("\n".join(md) + "\n", encoding="utf-8")
    (out / "pptx.json").write_text(json.dumps({"width": w, "height": h, "slides": slides}, ensure_ascii=False, indent=1), encoding="utf-8")
    done({"ok": True, "kind": "pptx", "slides": len(slides), "width": w, "height": h, "md": "PPTX.md"})


# ---- PDF (and Illustrator files saved with PDF content) -------------------------------------------------------------------------------------
PASSWORD = "has a password: use a copy without one"

def _open_pdf(file):
    """The document, or a clear answer for one that asks for a password (pdfium opens a file with only an owner password as it is; only a
    password needed to open it stops here). The password the owner gave in the chat comes in MW_PDF_PASSWORD, never on the command line."""
    try: import pypdfium2 as pdfium
    except Exception: fail("pypdfium2 is not installed (an app component: Settings > Health)")
    try: return pdfium.PdfDocument(file, password=os.environ.get("MW_PDF_PASSWORD") or None)
    except Exception as e:
        if "password" in str(e).lower(): done({"ok": False, "password": True, "error": PASSWORD})
        raise

# What Claude's Read takes in one go: about 25,000 tokens, 2,000 lines, 2,000 characters a line. A part stays well under all three, for
# text in any language (Vietnamese and Korean run near 2 characters a token).
TEXT_PART, TEXT_LINES, TEXT_LINE = 30000, 1500, 1800
SHEET_W, SHEET_MANY = 1568, 320          # the long edge Claude reads an image at, unscaled; past SHEET_MANY pages, 36 to a sheet

def _clean(t): return re.sub(r"[\x00-\x08\x0b\x0c\x0e-\x1f￾￿]", "", str(t).replace("\r\n", "\n").replace("\r", "\n")).strip()

def _text_parts(title, blocks, out):
    """Every block whole, across TEXT-NN.md files a single Read takes entire (a long line is wrapped, never cut)."""
    parts, cur, size, lines = [], [], 0, 0
    def flush():
        nonlocal cur, size, lines
        if cur: parts.append(cur)
        cur, size, lines = [], 0, 0
    for head, text in blocks:
        rows = [head, ""]
        for row in (text or "(no text: a picture page)").split("\n"):
            rows += [row[i:i + TEXT_LINE] for i in range(0, len(row), TEXT_LINE)] or [""]
        rows.append("")
        for row in rows:
            if size + len(row) + 1 > TEXT_PART or lines + 1 > TEXT_LINES: flush()
            cur.append(row); size += len(row) + 1; lines += 1
    flush()
    names = []
    for i, rows in enumerate(parts, 1):
        name = f"TEXT-{i:02d}.md"
        (out / name).write_text("\n".join([f"# {title}: its text, part {i} of {len(parts)}", ""] + rows) + "\n", encoding="utf-8", errors="replace")
        names.append(name)
    return names

def _sheets(items, out, per=16):
    """Pictures (path, label) on contact sheets, per to a sheet, each with its label; returns the sheets' names. The width is what Claude
    reads an image at, so a sheet costs its pixels / 750 tokens: about 1,800 for 16 pages of 16:9."""
    from PIL import Image, ImageDraw, ImageFont
    cols = 4 if per <= 16 else 6
    cell = SHEET_W // cols
    try: font = ImageFont.load_default(size=16 if cols == 4 else 13)
    except Exception: font = ImageFont.load_default()
    names = []
    for s in range(0, len(items), per):
        shots = []
        for path, label in items[s:s + per]:
            try:
                with Image.open(path) as im:
                    im.load()
                    # a picture with no ground is seen on mid grey, so a white logo and a black logo both show
                    if im.mode in ("RGBA", "LA", "P"):
                        rgba = im.convert("RGBA"); pic = Image.new("RGB", rgba.size, (128, 128, 128)); pic.paste(rgba, mask=rgba.getchannel("A"))
                    else: pic = im.convert("RGB")
            except Exception: pic = Image.new("RGB", (cell, cell // 2), (32, 32, 32))
            pic.thumbnail((cell - 4, round(cell * 1.42)), Image.LANCZOS)
            shots.append((pic, label))
        tall = max(p.height for p, _ in shots)
        rows = (len(shots) + cols - 1) // cols
        sheet = Image.new("RGB", (cell * cols, (tall + 4) * rows), (24, 24, 24))
        draw = ImageDraw.Draw(sheet)
        for i, (pic, label) in enumerate(shots):
            x, y = (i % cols) * cell + 2, (i // cols) * (tall + 4) + 2
            sheet.paste(pic, (x, y))
            box = draw.textbbox((x + 3, y + 2), label, font=font)
            draw.rectangle((box[0] - 3, box[1] - 2, box[2] + 3, box[3] + 2), fill=(0, 0, 0))
            draw.text((x + 3, y + 2), label, fill=(255, 255, 255), font=font)
        name = f"sheet-{s // per + 1:02d}.jpg"
        sheet.save(out / name, quality=80)
        names.append(name)
    return names

def pdf_whole(file, out, width, logo):
    """A guideline, whole: a page limit would drop what it says after page 40, and one picture a page would cost more than its text. Every
    page's text goes in TEXT-NN.md; every page is on a numbered contact sheet, to see its colours, type and layout; every page also waits
    as a picture (pages/), which costs nothing until Claude opens one for a closer look."""
    pdf = _open_pdf(file)
    n = len(pdf)
    (out / "pages").mkdir(exist_ok=True)
    blocks, items, chars = [], [], 0
    for i in range(n):
        page = pdf[i]
        w, _ = page.get_size()
        name = f"pages/page-{i + 1:03d}.jpg"
        page.render(scale=max(0.1, width / w)).to_pil().convert("RGB").save(out / name, quality=82)
        items.append((out / name, str(i + 1)))
        try: text = _clean(page.get_textpage().get_text_range())
        except Exception: text = ""
        chars += len(text)
        blocks.append((f"## Page {i + 1}", text))
    logo_file = None
    if logo and n:
        page = pdf[0]
        w, h = page.get_size()
        img = page.render(scale=max(0.1, 2000 / max(w, h)), fill_color=(0, 0, 0, 0)).to_pil().convert("RGBA")
        box = img.getchannel("A").getbbox()
        if box: img.crop(box).save(out / "logo.png"); logo_file = "logo.png"
    done({"ok": True, "kind": "pdf", "pages": n, "chars": chars, "text": _text_parts(Path(file).name, blocks, out),
          "sheets": _sheets(items, out, 16 if n <= SHEET_MANY else 36), "logo": logo_file})

def sheets_of(listing, out):
    """Pictures named one a line, on contact sheets labelled with their names."""
    paths = [l.strip() for l in Path(listing).read_text(encoding="utf-8").splitlines() if l.strip()]
    done({"ok": True, "kind": "sheet", "pictures": len(paths), "sheets": _sheets([(p, Path(p).name[:34]) for p in paths], out)})


# ---- Music -----------------------------------------------------------------------------------------------------------------------------------
def music_scan(file, out):
    import numpy as np, librosa
    y, sr = librosa.load(file, sr=22050, mono=True, duration=900)
    dur = len(y) / sr
    tempo, beats = librosa.beat.beat_track(y=y, sr=sr, units="time")
    tempo = float(np.atleast_1d(tempo)[0])
    rms = librosa.feature.rms(y=y, hop_length=512)[0]
    t = librosa.times_like(rms, sr=sr, hop_length=512)
    win = max(1, int(2 * sr / 512))          # loudest 2-second passages: where a film usually puts its big moments
    sm = np.convolve(rms, np.ones(win) / win, mode="same")
    peaks, used = [], []
    for idx in np.argsort(sm)[::-1]:
        tt = float(t[idx])
        if all(abs(tt - u) > 6 for u in used): used.append(tt); peaks.append(round(tt, 2))
        if len(peaks) >= 4: break
    info = {"ok": True, "kind": "music", "duration": round(dur, 2), "bpm": round(tempo, 1), "beats": len(beats),
            "first_beats": [round(float(b), 3) for b in beats[:16]], "bar_starts": [round(float(b), 3) for b in beats[::4][:24]], "loudest": sorted(peaks)}
    md = [f"# {Path(file).name}", "", f"{info['duration']} s, about {info['bpm']} BPM, {info['beats']} beats (librosa; a guide, the film's timing uses HyperFrames' beats).",
          "", f"Bar starts (every 4th beat, s): {', '.join(map(str, info['bar_starts']))}", "", f"Loudest passages (s): {', '.join(map(str, info['loudest']))}", ""]
    (out / "MUSIC.md").write_text("\n".join(md) + "\n", encoding="utf-8")
    done({**info, "md": "MUSIC.md"})


# ---- Archives --------------------------------------------------------------------------------------------------------------------------------
def zip_list(file, out):
    import zipfile
    with zipfile.ZipFile(file) as z:
        names = [n for n in z.namelist() if not n.endswith("/") and "__MACOSX" not in n]
    (out / "ZIP.md").write_text("\n".join([f"# {Path(file).name}", "", f"{len(names)} file(s):", ""] + [f"- {n}" for n in names[:500]]) + "\n", encoding="utf-8")
    done({"ok": True, "kind": "zip", "entries": len(names), "sample": names[:6], "md": "ZIP.md"})


if __name__ == "__main__":
    if len(sys.argv) < 4: fail(__doc__)
    cmd, file, out = sys.argv[1], sys.argv[2], Path(sys.argv[3])
    out.mkdir(parents=True, exist_ok=True)
    try:
        if cmd == "psd": psd_scan(file, out, int(arg("--width", 1600)))
        elif cmd == "psd-export": psd_export(file, out, arg("--frame", ""), int(arg("--width", 1920)))
        elif cmd == "pptx": pptx_scan(file, out)
        elif cmd == "pdf": pdf_whole(file, out, int(arg("--width", 1200)), "--logo" in sys.argv)
        elif cmd == "sheet": sheets_of(file, out)
        elif cmd == "music": music_scan(file, out)
        elif cmd == "zip": zip_list(file, out)
        else: fail(f"unknown command {cmd}")
    except SystemExit: raise
    except Exception as e: fail(f"{type(e).__name__}: {e}")
